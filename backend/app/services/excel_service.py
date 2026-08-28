from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook

from app.services.export_security import sanitize_cell_text
from app.services.results_export import MaterialTotal, ResultsExportDTO

DISCLAIMER_TEXT = (
    "MassIQ accelerates quantity takeoff and supports your review. "
    "It does not guarantee accuracy — confirm all quantities before use."
)


def build_excel_report(payload: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "MassIQ Report"
    rows = [
        ("Project", payload.get("project_name", "")),
        ("Plan", payload.get("plan_name", "")),
        ("Component", payload.get("component", "")),
        ("Detection Method", payload.get("detection_method", "")),
        ("Accepted Detection Area (m2)", payload.get("accepted_detection_area_m2", 0)),
        ("Added Correction Area (m2)", payload.get("added_correction_area_m2", 0)),
        ("Subtracted Correction Area (m2)", payload.get("subtracted_correction_area_m2", 0)),
        ("Final Area (m2)", payload.get("area_m2", 0)),
        ("Height (m)", payload.get("height_m", 0)),
        ("Volume (m3)", payload.get("volume_m3", 0)),
        ("Review Status", payload.get("review_status", "")),
        ("Notes", payload.get("notes", "")),
        ("Disclaimer", DISCLAIMER_TEXT),
    ]
    row_index = 0
    for row_index, (label, value) in enumerate(rows, start=1):
        ws.cell(row=row_index, column=1, value=label)
        ws.cell(row=row_index, column=2, value=value)

    deductions = payload.get("deductions") or []
    if deductions:
        header_row = row_index + 2
        ws.cell(row=header_row, column=1, value="Type")
        ws.cell(row=header_row, column=2, value="Count")
        ws.cell(row=header_row, column=3, value="Unit Area (m2)")
        ws.cell(row=header_row, column=4, value="Line Total (m2)")
        for offset, deduction in enumerate(deductions, start=1):
            data_row = header_row + offset
            deduction_type = str(deduction.get("type", "")).strip()
            label = deduction_type.capitalize() if deduction_type else "Deduction"
            ws.cell(row=data_row, column=1, value=label)
            ws.cell(row=data_row, column=2, value=deduction.get("count", 1))
            ws.cell(row=data_row, column=3, value=deduction.get("unit_area_m2", 0))
            ws.cell(row=data_row, column=4, value=deduction.get("line_total_m2", 0))

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


# -- R8: authoritative Results/Export workbook -----------------------------
# Deliberately a separate function, not a rewrite of build_excel_report
# above (R8 section 33 -- the legacy MVP export stays untouched). Formats
# already-computed ResultsExportDTO/MaterialTotal values; performs no
# quantity math of its own (R8 section 12).

AREA_FORMAT = "0.00"  # m^2, 2 decimal places (R8 section 17)
DIMENSION_FORMAT = "0.000"  # m, 3 decimal places
VOLUME_FORMAT = "0.000"  # m^3, 3 decimal places

QUANTITIES_HEADERS = (
    "Plan", "Page", "Material", "Material Code", "Area [m²]", "Dimension [m]",
    "Volume [m³]", "Status", "Calculation Version", "Quantity Result ID",
    "Detection Run ID", "Source Reference",
)
MATERIAL_TOTALS_HEADERS = ("Material", "Material Code", "Count", "Total Area [m²]", "Total Volume [m³]")
AUDIT_HEADERS = (
    "Quantity Result ID", "Detection Run ID", "Reference Type", "Reference ID",
    "Accepted Regions", "Rejected Regions", "Manual Additions", "Manual Subtractions",
    "Scale Method", "Calculation Version",
)


def _write_header_row(ws, row: int, headers: tuple[str, ...]) -> None:
    for col, header in enumerate(headers, start=1):
        ws.cell(row=row, column=col, value=header)


def _set_number(ws, row: int, col: int, value: float, number_format: str) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = number_format


def build_results_workbook(dto: ResultsExportDTO, material_totals: list[MaterialTotal]) -> bytes:
    """R8 sections 13/19: Summary + Quantities + Audit sheets, built
    entirely from an already-assembled ResultsExportDTO -- every
    area_m2/confirmed_dimension_m/volume_m3 value here is the exact
    persisted QuantityResult value (see results_export.build_export_dto),
    never recomputed."""
    wb = Workbook()

    # -- Summary --------------------------------------------------------
    summary = wb.active
    summary.title = "Summary"
    project_name = sanitize_cell_text(dto.project_name)
    summary_rows = [
        ("Project", project_name),
        ("Export Timestamp (UTC)", dto.export_timestamp.strftime("%Y-%m-%d %H:%M:%S")),
        ("Plan Count", dto.plan_count),
        ("Confirmed Quantity Count", len(dto.rows)),
        ("Disclaimer", DISCLAIMER_TEXT),
    ]
    for row_index, (label, value) in enumerate(summary_rows, start=1):
        summary.cell(row=row_index, column=1, value=label)
        summary.cell(row=row_index, column=2, value=value)

    # Grouped material totals (R8 sections 14-15) -- never a single
    # cross-material grand total.
    totals_header_row = len(summary_rows) + 3
    summary.cell(row=totals_header_row - 1, column=1, value="Totals by Material")
    _write_header_row(summary, totals_header_row, MATERIAL_TOTALS_HEADERS)
    for offset, total in enumerate(material_totals, start=1):
        row = totals_header_row + offset
        summary.cell(row=row, column=1, value=sanitize_cell_text(total.material_name))
        summary.cell(row=row, column=2, value=sanitize_cell_text(total.material_code))
        summary.cell(row=row, column=3, value=total.count)
        _set_number(summary, row, 4, total.total_area_m2, AREA_FORMAT)
        _set_number(summary, row, 5, total.total_volume_m3, VOLUME_FORMAT)

    # -- Quantities -------------------------------------------------------
    quantities = wb.create_sheet("Quantities")
    _write_header_row(quantities, 1, QUANTITIES_HEADERS)
    for offset, row in enumerate(dto.rows, start=1):
        r = offset + 1
        quantities.cell(row=r, column=1, value=sanitize_cell_text(row.plan_name))
        quantities.cell(row=r, column=2, value=row.page_number)
        quantities.cell(row=r, column=3, value=sanitize_cell_text(row.material_name))
        quantities.cell(row=r, column=4, value=sanitize_cell_text(row.material_code))
        _set_number(quantities, r, 5, row.area_m2, AREA_FORMAT)
        _set_number(quantities, r, 6, row.confirmed_dimension_m, DIMENSION_FORMAT)
        _set_number(quantities, r, 7, row.volume_m3, VOLUME_FORMAT)
        quantities.cell(row=r, column=8, value=row.status.value)
        quantities.cell(row=r, column=9, value=row.calculation_version)
        quantities.cell(row=r, column=10, value=str(row.quantity_result_id))
        quantities.cell(row=r, column=11, value=str(row.detection_run_id))
        quantities.cell(row=r, column=12, value=f"{row.reference_type}:{row.reference_id}")

    # -- Audit (R8 section 19) -----------------------------------------
    audit = wb.create_sheet("Audit")
    _write_header_row(audit, 1, AUDIT_HEADERS)
    for offset, row in enumerate(dto.rows, start=1):
        r = offset + 1
        audit.cell(row=r, column=1, value=str(row.quantity_result_id))
        audit.cell(row=r, column=2, value=str(row.detection_run_id))
        audit.cell(row=r, column=3, value=row.reference_type)
        audit.cell(row=r, column=4, value=str(row.reference_id))
        audit.cell(row=r, column=5, value=row.accepted_region_count)
        audit.cell(row=r, column=6, value=row.rejected_region_count)
        audit.cell(row=r, column=7, value=row.manual_add_count)
        audit.cell(row=r, column=8, value=row.manual_subtract_count)
        audit.cell(row=r, column=9, value=row.scale_method or "")
        audit.cell(row=r, column=10, value=row.calculation_version)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
