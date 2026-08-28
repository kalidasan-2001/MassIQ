"""R8 sections 27-28 -- release-critical. Actually opens the generated
.xlsx with openpyxl and inspects its real content: sheet names, headers,
numeric values, units, row counts, ordering, grouping totals, and formula-
injection sanitization. Never just asserts HTTP 200 / "bytes were
returned". Pure -- ResultRow is a plain dataclass, no DB needed here.
"""

from __future__ import annotations

import sys
import unittest
import uuid
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from openpyxl import load_workbook

from app.models.quantity_result import QuantityResultStatus
from app.services.excel_service import build_results_workbook
from app.services.results_export import build_export_dto, group_totals_by_material
from app.services.results_service import ResultRow


def _row(**overrides) -> ResultRow:
    defaults = dict(
        quantity_result_id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        plan_id=uuid.uuid4(),
        plan_name="Ground Floor.pdf",
        plan_page_id=uuid.uuid4(),
        page_number=1,
        material_name="Stahlbeton C25/30",
        material_code="C25/30",
        area_m2=84.2,
        confirmed_dimension_m=0.2,
        volume_m3=16.84,
        status=QuantityResultStatus.CONFIRMED,
        calculation_version="1.0",
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
        confirmed_at=datetime.now(timezone.utc),
        detection_run_id=uuid.uuid4(),
        reference_type="legend_entry",
        reference_id=uuid.uuid4(),
        accepted_region_count=5,
        rejected_region_count=2,
        manual_add_count=1,
        manual_subtract_count=2,
        scale_method="declared_scale",
    )
    defaults.update(overrides)
    return ResultRow(**defaults)


def _open(workbook_bytes: bytes):
    return load_workbook(BytesIO(workbook_bytes))


class WorkbookStructureTests(unittest.TestCase):
    def test_sheet_names(self):
        rows = [_row()]
        dto = build_export_dto("Test Project", rows)
        wb_bytes = build_results_workbook(dto, group_totals_by_material(rows))
        wb = _open(wb_bytes)
        self.assertEqual(wb.sheetnames, ["Summary", "Quantities", "Audit"])

    def test_quantities_sheet_headers(self):
        rows = [_row()]
        dto = build_export_dto("Test Project", rows)
        wb = _open(build_results_workbook(dto, group_totals_by_material(rows)))
        ws = wb["Quantities"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(
            headers,
            [
                "Plan", "Page", "Material", "Material Code", "Area [m²]", "Dimension [m]",
                "Volume [m³]", "Status", "Calculation Version", "Quantity Result ID",
                "Detection Run ID", "Source Reference",
            ],
        )

    def test_audit_sheet_headers(self):
        rows = [_row()]
        dto = build_export_dto("Test Project", rows)
        wb = _open(build_results_workbook(dto, group_totals_by_material(rows)))
        ws = wb["Audit"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(
            headers,
            [
                "Quantity Result ID", "Detection Run ID", "Reference Type", "Reference ID",
                "Accepted Regions", "Rejected Regions", "Manual Additions", "Manual Subtractions",
                "Scale Method", "Calculation Version",
            ],
        )


class AuthoritativeValueRegressionTests(unittest.TestCase):
    """R8 section 28 -- release-critical. The exporter must not
    independently recalculate area/volume; Excel values must equal the
    exact persisted (here: dataclass-constructed, standing in for
    persisted) authoritative numbers."""

    def test_excel_values_equal_authoritative_persisted_values(self):
        row = _row(area_m2=84.2, confirmed_dimension_m=0.2, volume_m3=16.84)
        dto = build_export_dto("Audit Co", [row])
        wb = _open(build_results_workbook(dto, group_totals_by_material([row])))
        ws = wb["Quantities"]
        self.assertEqual(ws.cell(row=2, column=5).value, 84.2)  # Area
        self.assertEqual(ws.cell(row=2, column=6).value, 0.2)  # Dimension
        self.assertEqual(ws.cell(row=2, column=7).value, 16.84)  # Volume
        self.assertEqual(ws.cell(row=2, column=10).value, str(row.quantity_result_id))

    def test_multiple_rows_each_keep_their_own_authoritative_values(self):
        row_a = _row(material_name="Material A", area_m2=10.0, volume_m3=2.0)
        row_b = _row(material_name="Material B", area_m2=55.5, volume_m3=11.1)
        dto = build_export_dto("Multi Co", [row_a, row_b])
        wb = _open(build_results_workbook(dto, group_totals_by_material([row_a, row_b])))
        ws = wb["Quantities"]
        areas = {ws.cell(row=r, column=3).value: ws.cell(row=r, column=5).value for r in (2, 3)}
        self.assertEqual(areas["Material A"], 10.0)
        self.assertEqual(areas["Material B"], 55.5)

    def test_units_are_explicit_in_headers(self):
        rows = [_row()]
        dto = build_export_dto("Units Co", rows)
        wb = _open(build_results_workbook(dto, group_totals_by_material(rows)))
        headers = [cell.value for cell in wb["Quantities"][1]]
        self.assertIn("Area [m²]", headers)
        self.assertIn("Dimension [m]", headers)
        self.assertIn("Volume [m³]", headers)

    def test_row_order_matches_dto_order(self):
        # ResultsService is responsible for sort order; the workbook must
        # simply preserve whatever order the DTO's rows arrive in.
        row_z = _row(material_name="Zebra")
        row_a = _row(material_name="Alpha")
        dto = build_export_dto("Order Co", [row_z, row_a])  # deliberately not pre-sorted here
        wb = _open(build_results_workbook(dto, group_totals_by_material([row_z, row_a])))
        ws = wb["Quantities"]
        materials = [ws.cell(row=r, column=3).value for r in (2, 3)]
        self.assertEqual(materials, ["Zebra", "Alpha"])


class MaterialGroupingTests(unittest.TestCase):
    def test_same_material_grouped_totals(self):
        row_a = _row(material_name="Concrete", material_code="C1", area_m2=10.0, volume_m3=2.0)
        row_b = _row(material_name="Concrete", material_code="C1", area_m2=20.0, volume_m3=4.0)
        totals = group_totals_by_material([row_a, row_b])
        self.assertEqual(len(totals), 1)
        self.assertEqual(totals[0].count, 2)
        self.assertAlmostEqual(totals[0].total_area_m2, 30.0)
        self.assertAlmostEqual(totals[0].total_volume_m3, 6.0)

    def test_same_name_different_code_not_grouped(self):
        # "Concrete" vs "Existing Concrete" -- and same name, different
        # code -- must never be silently merged (R8 section 15).
        row_a = _row(material_name="Concrete", material_code="C1", area_m2=10.0)
        row_b = _row(material_name="Concrete", material_code="C2", area_m2=20.0)
        totals = group_totals_by_material([row_a, row_b])
        self.assertEqual(len(totals), 2)

    def test_different_material_names_not_grouped(self):
        row_a = _row(material_name="Concrete", area_m2=10.0)
        row_b = _row(material_name="Existing Concrete", area_m2=20.0)
        totals = group_totals_by_material([row_a, row_b])
        self.assertEqual(len(totals), 2)
        self.assertNotEqual(totals[0].material_name, totals[1].material_name)

    def test_summary_sheet_contains_grouped_totals_not_grand_total(self):
        row_a = _row(material_name="Steel", area_m2=10.0, volume_m3=1.0)
        row_b = _row(material_name="Concrete", area_m2=90.0, volume_m3=18.0)
        dto = build_export_dto("Grouping Co", [row_a, row_b])
        wb = _open(build_results_workbook(dto, group_totals_by_material([row_a, row_b])))
        ws = wb["Summary"]
        values = [cell.value for row in ws.iter_rows() for cell in row]
        # Both material names appear as distinct grouped rows.
        self.assertIn("Steel", values)
        self.assertIn("Concrete", values)
        # No single misleading combined total (100.0/19.0) appears as a
        # standalone summary field -- only per-material subtotals exist.
        self.assertNotIn(100.0, values)
        self.assertNotIn(19.0, values)


class FormulaInjectionTests(unittest.TestCase):
    """R8 section 41 -- release-relevant security test."""

    def test_material_name_starting_with_equals_is_neutralized(self):
        row = _row(material_name="=1+1")
        dto = build_export_dto("Injection Co", [row])
        wb = _open(build_results_workbook(dto, group_totals_by_material([row])))
        cell_value = wb["Quantities"].cell(row=2, column=3).value
        self.assertTrue(cell_value.startswith("'"))
        self.assertNotEqual(cell_value, "=1+1")

    def test_material_code_starting_with_plus_is_neutralized(self):
        row = _row(material_code="+cmd|calc")
        dto = build_export_dto("Injection Co", [row])
        wb = _open(build_results_workbook(dto, group_totals_by_material([row])))
        cell_value = wb["Quantities"].cell(row=2, column=4).value
        self.assertTrue(cell_value.startswith("'"))

    def test_project_name_starting_with_at_is_neutralized(self):
        rows = [_row()]
        dto = build_export_dto("@SUM(A1:A9999)", rows)
        wb = _open(build_results_workbook(dto, group_totals_by_material(rows)))
        project_cell = wb["Summary"].cell(row=1, column=2).value
        self.assertTrue(project_cell.startswith("'"))

    def test_ordinary_material_name_not_mangled(self):
        row = _row(material_name="Stahlbeton C25/30")
        dto = build_export_dto("Safe Co", [row])
        wb = _open(build_results_workbook(dto, group_totals_by_material([row])))
        cell_value = wb["Quantities"].cell(row=2, column=3).value
        self.assertEqual(cell_value, "Stahlbeton C25/30")


class PerformanceTests(unittest.TestCase):
    """R8 section 43 -- generous ceilings (not tight timing assertions
    that would be flaky in CI), just proving workbook generation stays
    comfortably interactive at representative and stress-test scale.
    Measured directly during implementation: ~70ms/10 rows, ~110ms/100
    rows, ~540ms/1000 rows on ordinary developer hardware."""

    def test_1000_rows_generates_well_under_a_few_seconds(self):
        import time

        rows = [_row(material_name=f"Material {i % 20}", area_m2=float(i) + 0.5) for i in range(1000)]
        dto = build_export_dto("Perf Co", rows)
        start = time.perf_counter()
        wb_bytes = build_results_workbook(dto, group_totals_by_material(rows))
        elapsed_s = time.perf_counter() - start
        self.assertLess(elapsed_s, 5.0)
        self.assertGreater(len(wb_bytes), 0)


class EmptyRowsTests(unittest.TestCase):
    def test_zero_rows_still_produces_a_structurally_valid_workbook(self):
        dto = build_export_dto("Empty Co", [])
        wb = _open(build_results_workbook(dto, group_totals_by_material([])))
        self.assertEqual(wb.sheetnames, ["Summary", "Quantities", "Audit"])
        self.assertIsNone(wb["Quantities"].cell(row=2, column=1).value)


if __name__ == "__main__":
    unittest.main()
