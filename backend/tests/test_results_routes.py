"""API-level integration tests against a real Postgres database. Requires
the docker-compose Postgres service to be running:
    docker compose up -d postgres
"""

from __future__ import annotations

import sys
import tempfile
import unittest
import uuid
from io import BytesIO
from pathlib import Path

import cv2
import numpy as np
from fastapi.testclient import TestClient
from openpyxl import load_workbook

BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

import app.routes.results as results_route_module
from tests import hatch_fixtures as hfx
from tests import pdf_fixtures as fx
from tests.db_test_support import build_test_engine, new_sessionmaker, truncate_all

from app.db.database import get_db
from app.main import app
from app.models.detected_region import DetectedRegion, DetectedRegionStatus
from app.schemas.legend_entry import LegendEntryUpdate
from app.schemas.project import ProjectCreate
from app.services.detection_service import DetectionService
from app.services.hatch_feature_service import HatchFeatureService
from app.services.legend_crop_service import RegionSelection
from app.services.legend_service import LegendService
from app.services.ocr_service import OcrResult, OcrService
from app.services.plan_scale_service import PlanScaleService
from app.services.plan_service import PlanService
from app.services.project_service import ProjectService
from app.services.quantity_service import QuantityService
from app.services.results_service import ResultsService
from app.services.storage_service import StorageService


class _FakeOcrProvider:
    name = "fake"

    def extract_text(self, image_path):
        return OcrResult(text="Stahlbeton C25/30", provider=self.name, confidence=0.9)


def _encode_png(image) -> bytes:
    ok, buffer = cv2.imencode(".png", image)
    assert ok
    return buffer.tobytes()


class ResultsRoutesTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = build_test_engine()
        cls.Session = new_sessionmaker(cls.engine)

    @classmethod
    def tearDownClass(cls):
        cls.engine.dispose()

    def setUp(self):
        truncate_all(self.engine)
        self.session = self.Session()
        self._tmp = tempfile.TemporaryDirectory()
        self.storage = StorageService(root=Path(self._tmp.name))

        def _override_get_db():
            yield self.session

        app.dependency_overrides[get_db] = _override_get_db
        app.dependency_overrides[results_route_module.get_results_service] = lambda: ResultsService(self.session)

        self.client = TestClient(app)

        self.project_service = ProjectService(self.session)
        self.plan_service = PlanService(self.session, storage=self.storage)
        self.legend_service = LegendService(self.session, storage=self.storage, ocr=OcrService(provider=_FakeOcrProvider()))
        self.feature_service = HatchFeatureService(self.session, storage=self.storage)
        self.detection_service = DetectionService(self.session, storage=self.storage)
        self.scale_service = PlanScaleService(self.session, storage=self.storage)
        self.quantity_service = QuantityService(self.session, storage=self.storage)

        self.project = self.project_service.create_project(ProjectCreate(name="Route Results Project"))
        self.plan = self.plan_service.upload_plan(self.project.id, "vector.pdf", fx.build_vector_pdf_bytes())

    def tearDown(self):
        app.dependency_overrides.clear()
        self.session.close()
        self._tmp.cleanup()

    def _make_result(self, material_name="Stahlbeton C25/30", confirm=True, area_w=0.1, dimension_m=2.0):
        entry = self.legend_service.create_draft(self.project.id, self.plan.id, page_number=1)
        entry = self.legend_service.save_pattern_selection(
            self.project.id, self.plan.id, entry.id, RegionSelection(x=0.05, y=0.05, width=0.2, height=0.1)
        )
        self.legend_service.save_description_selection(
            self.project.id, self.plan.id, entry.id, RegionSelection(x=0.5, y=0.5, width=0.2, height=0.1)
        )
        self.legend_service.update_entry(
            self.project.id, self.plan.id, entry.id,
            LegendEntryUpdate(corrected_text=material_name, material_name=material_name),
        )
        confirmed_entry = self.legend_service.confirm(self.project.id, self.plan.id, entry.id)
        crop_path = self.storage.resolve_legend_crop(confirmed_entry.pattern_image_reference)
        crop_path.write_bytes(_encode_png(hfx.family_a_parallel_45()))
        self.feature_service.compute_features(self.project.id, self.plan.id, confirmed_entry.id)
        page = self.plan_service.get_page(self.project.id, self.plan.id, 1)
        preview_path = self.storage.resolve_preview(page.preview_reference)
        preview_path.write_bytes(_encode_png(np.full((700, 900, 3), 255, dtype=np.uint8)))
        run = self.detection_service.start_run(self.project.id, self.plan.id, 1, legend_entry_id=confirmed_entry.id)
        self.session.add(
            DetectedRegion(
                detection_run_id=run.id, x=0, y=0, width=area_w, height=0.1,
                similarity=0.9, evidence_coverage=0.9, tile_count=1, status=DetectedRegionStatus.ACCEPTED,
            )
        )
        self.session.commit()
        if self.scale_service.get_scale_or_none(self.project.id, self.plan.id, 1) is None:
            self.scale_service.confirm_calibrated_distance(
                self.project.id, self.plan.id, 1, calibrated_distance_plan_points=1.0, calibrated_distance_real_m=1.0
            )
        result = self.quantity_service.calculate(self.project.id, self.plan.id, run.id, confirmed_dimension_m=dimension_m)
        if confirm:
            result = self.quantity_service.confirm(self.project.id, self.plan.id, run.id)
        return result, run

    # -- List / detail -----------------------------------------------------

    def test_list_results_empty_project(self):
        response = self.client.get(f"/api/projects/{self.project.id}/results")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), [])

    def test_list_results_includes_draft_and_confirmed(self):
        self._make_result(material_name="Confirmed Mat", confirm=True)
        self._make_result(material_name="Draft Mat", confirm=False)
        response = self.client.get(f"/api/projects/{self.project.id}/results")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(len(response.json()), 2)
        statuses = {row["status"] for row in response.json()}
        self.assertEqual(statuses, {"confirmed", "draft"})

    def test_get_result_detail(self):
        result, _run = self._make_result()
        response = self.client.get(f"/api/projects/{self.project.id}/results/{result.id}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["quantity_result_id"], str(result.id))

    def test_get_unknown_result_is_404(self):
        response = self.client.get(f"/api/projects/{self.project.id}/results/{uuid.uuid4()}")
        self.assertEqual(response.status_code, 404)

    def test_unknown_project_is_404(self):
        response = self.client.get(f"/api/projects/{uuid.uuid4()}/results")
        self.assertEqual(response.status_code, 404)

    # -- Cross-project isolation (R8 section 30) ----------------------------

    def test_cross_project_result_detail_is_404(self):
        result, _run = self._make_result()
        other_project = self.project_service.create_project(ProjectCreate(name="Other"))
        response = self.client.get(f"/api/projects/{other_project.id}/results/{result.id}")
        self.assertEqual(response.status_code, 404)

    def test_cross_project_listing_never_leaks(self):
        self._make_result(material_name="Project A Material")
        other_project = self.project_service.create_project(ProjectCreate(name="Other"))
        response = self.client.get(f"/api/projects/{other_project.id}/results")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), [])

    def test_cross_project_export_never_leaks(self):
        self._make_result(material_name="Project A Material")
        other_project = self.project_service.create_project(ProjectCreate(name="Other"))
        # Other project has no confirmed results of its own -> 409, not
        # Project A's data.
        response = self.client.post(f"/api/projects/{other_project.id}/results/export")
        self.assertEqual(response.status_code, 409)

    # -- Empty export (R8 sections 9/32) -------------------------------------

    def test_export_with_no_confirmed_results_is_409(self):
        response = self.client.post(f"/api/projects/{self.project.id}/results/export")
        self.assertEqual(response.status_code, 409)

    def test_export_preflight_reflects_confirmed_count(self):
        response = self.client.get(f"/api/projects/{self.project.id}/results/export/preflight")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"confirmed_count": 0, "can_export": False})
        self._make_result()
        response = self.client.get(f"/api/projects/{self.project.id}/results/export/preflight")
        self.assertEqual(response.json(), {"confirmed_count": 1, "can_export": True})

    # -- Draft exclusion from export (R8 section 31) -------------------------

    def test_export_excludes_draft_rows(self):
        self._make_result(material_name="Confirmed One", confirm=True)
        self._make_result(material_name="Confirmed Two", confirm=True)
        self._make_result(material_name="Draft Three", confirm=False)

        response = self.client.post(f"/api/projects/{self.project.id}/results/export")
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(BytesIO(response.content))
        ws = wb["Quantities"]
        data_rows = [row for row in ws.iter_rows(min_row=2) if row[0].value is not None]
        self.assertEqual(len(data_rows), 2)
        materials = {row[2].value for row in data_rows}
        self.assertEqual(materials, {"Confirmed One", "Confirmed Two"})
        self.assertNotIn("Draft Three", materials)

    # -- Filename (R8 section 21/42) -----------------------------------------

    def test_export_filename_is_sanitized_and_dated(self):
        self._make_result()
        response = self.client.post(f"/api/projects/{self.project.id}/results/export")
        self.assertEqual(response.status_code, 200)
        disposition = response.headers["content-disposition"]
        # Spaces are legal filename characters and stay as-is; only unsafe
        # characters get stripped (see test_export_security.py).
        self.assertIn("MassIQ_Route Results Project_quantities_", disposition)
        self.assertTrue(disposition.strip().endswith('.xlsx"'))
        self.assertNotIn("/", disposition)
        self.assertNotIn("..", disposition)

    # -- Authoritative-value regression through the real HTTP path ----------

    def test_exported_values_match_authoritative_quantity_result(self):
        result, _run = self._make_result(dimension_m=3.0)
        response = self.client.post(f"/api/projects/{self.project.id}/results/export")
        wb = load_workbook(BytesIO(response.content))
        ws = wb["Quantities"]
        self.assertAlmostEqual(ws.cell(row=2, column=5).value, result.final_area_m2, places=6)
        self.assertAlmostEqual(ws.cell(row=2, column=6).value, result.confirmed_dimension_m, places=6)
        self.assertAlmostEqual(ws.cell(row=2, column=7).value, result.volume_m3, places=6)

    # -- Quantity change / re-export (R8 section 29) -------------------------

    def test_reexport_after_recalculation_reflects_new_state_not_stale_cache(self):
        result, run = self._make_result(area_w=0.1, dimension_m=1.0, confirm=True)
        first_export = self.client.post(f"/api/projects/{self.project.id}/results/export")
        wb_first = load_workbook(BytesIO(first_export.content))
        area_first = wb_first["Quantities"].cell(row=2, column=5).value

        # User adds another accepted region and recalculates.
        self.session.add(
            DetectedRegion(
                detection_run_id=run.id, x=0.5, y=0.5, width=0.1, height=0.1,
                similarity=0.9, evidence_coverage=0.9, tile_count=1, status=DetectedRegionStatus.ACCEPTED,
            )
        )
        self.session.commit()
        recalculated = self.quantity_service.calculate(self.project.id, self.plan.id, run.id, confirmed_dimension_m=1.0)
        self.quantity_service.confirm(self.project.id, self.plan.id, run.id)

        second_export = self.client.post(f"/api/projects/{self.project.id}/results/export")
        wb_second = load_workbook(BytesIO(second_export.content))
        area_second = wb_second["Quantities"].cell(row=2, column=5).value

        self.assertGreater(area_second, area_first)
        self.assertAlmostEqual(area_second, recalculated.final_area_m2, places=6)


if __name__ == "__main__":
    unittest.main()
